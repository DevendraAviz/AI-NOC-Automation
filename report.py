"""Excel report: the same matrix as Dev's sheet, plus details and a summary."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
HEAD = Font(name="Arial", size=11, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="B4A7D6")            # Dev's sheet header colour
FILL = {"PASS": "C6EFCE", "FAIL": "FFC7CE", "NA": "D9D9D9", "BLOCKED": "FFEB9C", "XFAIL": "F4CCCC"}
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
EDGE = Side(style="thin", color="000000")
BOX = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)
ORDER = ["FAIL", "XFAIL", "BLOCKED", "NA", "PASS"]               # worst first when merging scopes


def _cell(ws, row, col, value, font=FONT, align=WRAP, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font, c.alignment, c.border = font, align, BOX
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c


def write_report(results: list[dict], prompts: list, connectors: list, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"NCP_MCP_Prompt_Results_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    wb = Workbook()

    # ---- Matrix: Prompt | <connector titles> | Comments -------------------------------------
    ws = wb.active
    ws.title = "Matrix"
    headers = ["Prompt"] + [c.title for c in connectors] + ["Comments"]
    for i, h in enumerate(headers, 1):
        _cell(ws, 1, i, h, HEAD, CENTER, "B4A7D6")
    ws.column_dimensions["A"].width = 62
    for i in range(2, len(headers)):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 22
    ws.column_dimensions[ws.cell(row=1, column=len(headers)).column_letter].width = 80
    by_key = {}
    for r in results:
        by_key.setdefault((r["id"], r["connector"]), []).append(r)
    for row_i, p in enumerate(prompts, 2):
        _cell(ws, row_i, 1, p.prompt)
        comments = []
        for col_i, conn in enumerate(connectors, 2):
            runs = by_key.get((p.id, conn.key), [])
            if not runs:
                _cell(ws, row_i, col_i, "", align=CENTER)
                continue
            status = min((r["status"] for r in runs), key=ORDER.index)
            _cell(ws, row_i, col_i, status, BOLD, CENTER, FILL.get(status))
            for r in runs:
                if r["status"] != "PASS":
                    scope = f" [{r['scope']}]" if len(runs) > 1 else ""
                    comments.append(f"{conn.title.split(' (')[0]}{scope}: {r['status']} — {r['reason']}")
        _cell(ws, row_i, len(headers), "\n".join(comments))
    ws.freeze_panes = "B2"

    # ---- Details -------------------------------------------------------------------------------
    d = wb.create_sheet("Details")
    cols = ["ID", "Connector", "Scope", "Result", "Reason", "Expected (source)", "Prompt sent", "Device",
            "Follow-ups", "Seconds", "Conversation", "Judge note", "NCP answer"]
    widths = [6, 24, 8, 9, 60, 60, 50, 18, 40, 8, 12, 40, 90]
    for i, (h, w) in enumerate(zip(cols, widths), 1):
        _cell(d, 1, i, h, HEAD, CENTER, "B4A7D6")
        d.column_dimensions[d.cell(row=1, column=i).column_letter].width = w
    for i, r in enumerate(results, 2):
        values = [r["id"], r["title"], r["scope"], r["status"], r["reason"], r["expected"], r["sent"],
                  r["device"], " || ".join(f"{a} -> {b[:150]}" for a, b in r["followups"]), r["seconds"],
                  r["conversation_id"], r["judge"], (r["answer"] or "")[:32000]]
        for j, v in enumerate(values, 1):
            _cell(d, i, j, v, fill=FILL.get(r["status"]) if j == 4 else None)
    d.freeze_panes = "C2"
    d.auto_filter.ref = f"A1:M{max(2, len(results) + 1)}"

    # ---- Summary (live formulas over Details) --------------------------------------------------
    s = wb.create_sheet("Summary", 0)
    _cell(s, 1, 1, "Connector", HEAD, CENTER, "B4A7D6")
    for j, st in enumerate(["PASS", "FAIL", "NA", "BLOCKED", "XFAIL", "Total"], 2):
        _cell(s, 1, j, st, HEAD, CENTER, "B4A7D6")
    last = len(results) + 1
    for i, conn in enumerate(connectors, 2):
        _cell(s, i, 1, conn.title, BOLD)
        for j, st in enumerate(["PASS", "FAIL", "NA", "BLOCKED", "XFAIL"], 2):
            _cell(s, i, j, f'=COUNTIFS(Details!$B$2:$B${last},$A{i},Details!$D$2:$D${last},"{st}")',
                  align=CENTER)
        _cell(s, i, 7, f"=SUM(B{i}:F{i})", BOLD, CENTER)
    s.column_dimensions["A"].width = 30
    note_row = len(connectors) + 3
    s.cell(row=note_row, column=1, value=(
        "PASS = matches the source · FAIL = wrong/missing/invented · NA = the product has no such data "
        "and NCP said so · BLOCKED = ground truth could not be read (not an NCP result) · "
        "XFAIL = known issue from the prompt sheet")).font = FONT
    wb.save(path)
    return path
