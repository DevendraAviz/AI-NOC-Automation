"""Offline self-tests: the grading logic, the chat client and the report, with fake data.
No network. Run with plain `pytest` (pytest.ini points here)."""
from __future__ import annotations

import asyncio
import json
import socket
import threading
import time
from dataclasses import replace

import pytest
import websockets
from openpyxl import load_workbook

from ncp_suite import settings
from ncp_suite.chat import NcpChat
from ncp_suite.chat.client import error_kind
from ncp_suite.chat.policy import followup_reply, is_followup
from ncp_suite.chat.stream import AnswerStream, inline_ui
from ncp_suite.grading.checks import CHECKS, Ctx, evaluate
from ncp_suite.prompts import PromptRow, load_prompts
from ncp_suite.reporting.excel import write_report
from ncp_suite.results import PromptResult
from ncp_suite.runner import run_case
from ncp_suite.settings import ChatSettings
from ncp_suite.truth.base import Component, Device, Interface, Link, NoTruth, Source

pytestmark = pytest.mark.offline


@pytest.fixture(autouse=True)
def _no_llm_reader(monkeypatch):
    """Self-tests never call the LLM reader on the network (a test that wants it fakes it)."""
    monkeypatch.setattr(settings, "EXTRACT_URL", "")


# ---------------------------------------------------------------- fake source
class FakeSource(Source):
    def __init__(self, links_unsupported=False, metrics_broken=False, no_mem=False):
        super().__init__(settings.Connector("fake", "Fake (Local MCP)", "#fake", "x:y", "http://x", "u", "p"))
        self.EMPTY_MEANS_UNSUPPORTED = ({"links"} if links_unsupported else set()) | ({"mem"} if no_mem else set())
        self.links_unsupported, self.metrics_broken, self.no_mem = links_unsupported, metrics_broken, no_mem

    def _devices(self):
        return [Device("leaf-1", "10.0.0.1", "N9K-C93180YC-FX", "FDO111", "10.3(3)", healthy=True),
                Device("leaf-2", "10.0.0.2", "N9K-C93180YC-FX", "FDO222", "10.3(3)", healthy=False, reason="down"),
                Device("spine-1", "10.0.0.3", "N9K-C9332C", "FDO333", "10.2(5)", healthy=True)]

    def _metrics(self):
        if self.metrics_broken:
            raise NoTruth("HTTP 404 on metrics")
        data = {"leaf-1": {"cpu": 12, "mem": 40, "temp": 35}, "leaf-2": {"cpu": 91, "mem": 80, "temp": 52},
                "spine-1": {"cpu": 45, "mem": 60, "temp": 41}}
        if self.no_mem:
            for v in data.values():
                v["mem"] = None
        return data

    def _interfaces(self, dev):
        return [Interface(dev.name, "Ethernet1/1", "up"), Interface(dev.name, "Ethernet1/2", "down"),
                Interface(dev.name, "Ethernet1/3", "up")]

    def _links(self):
        if self.links_unsupported:
            return []
        return [Link("leaf-1", "Eth1/49", "spine-1", "Eth1/1"), Link("leaf-2", "Eth1/49", "spine-1", "Eth1/2")]

    def _components(self):
        return [Component("leaf-1", "fan", "Fan1", "ok", True), Component("leaf-1", "psu", "PSU1", "ok", True),
                Component("leaf-2", "psu", "PSU2", "failed", False), Component("spine-1", "fan", "Fan1", "ok", True)]


def T(header, *rows):
    return "\n".join(["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
                     + ["| " + " | ".join(map(str, r)) + " |" for r in rows])


GOOD = {
    "devices_list": T(["Hostname", "IP"], ["leaf-1", "10.0.0.1"], ["leaf-2", "10.0.0.2"], ["spine-1", "10.0.0.3"]),
    "devices_count": "There are 3 devices in your network.",
    "devices_fields": T(["Hostname", "Mgmt IP", "Model", "Serial", "Version"],
                        ["leaf-1", "10.0.0.1", "N9K-C93180YC-FX", "FDO111", "10.3(3)"],
                        ["leaf-2", "10.0.0.2", "N9K-C93180YC-FX", "FDO222", "10.3(3)"],
                        ["spine-1", "10.0.0.3", "N9K-C9332C", "FDO333", "10.2(5)"]),
    "os_version_counts": T(["OS version", "Devices"], ["10.3(3)", 2], ["10.2(5)", 1]),
    "models_list": "Models in use: N9K-C93180YC-FX and N9K-C9332C.",
    "unhealthy_devices": "Only leaf-2 looks unhealthy: it is down / unreachable.",
    "cpu_all": T(["Device", "CPU %"], ["leaf-1", 13], ["leaf-2", 89.5], ["spine-1", 44]),
    "mem_all": T(["Device", "Memory utilization"], ["leaf-1", "41%"], ["leaf-2", "79%"], ["spine-1", "60%"]),
    "cpu_mem_device": "Device leaf-1: CPU utilization 13%, memory utilization 41%.",
    "cpu_top": "leaf-2 has the highest CPU at 91%, followed by spine-1 (45%).",
    "cpu_above": T(["Device", "CPU"], ["leaf-2", "91%"]),
    "mem_above": T(["Device", "Memory"], ["leaf-2", "80%"]),
    "interfaces_list": "leaf-1 interfaces: Ethernet1/1, Ethernet1/2, Ethernet1/3.",
    "interfaces_down": "On leaf-1, Eth1/2 is down.",
    "interface_counters": T(["Interface", "RX bytes", "TX bytes", "Errors"],
                            ["Ethernet1/1", 100, 200, 0], ["Ethernet1/2", 0, 0, 0], ["Ethernet1/3", 5, 6, 0]),
    "links": T(["Local", "Remote"], ["leaf-1 Eth1/49", "spine-1 Eth1/1"], ["leaf-2 Eth1/49", "spine-1 Eth1/2"]),
    "fan_psu": T(["Device", "Part", "Status"], ["leaf-1", "Fan1", "OK"], ["leaf-1", "PSU1", "OK"],
                 ["leaf-2", "PSU2", "Failed"], ["spine-1", "Fan1", "OK"]),
    "temperature": T(["Device", "Temperature (°C)"], ["leaf-1", 36], ["leaf-2", 51], ["spine-1", 41]),
    "health_summary": T(["Device", "Health"], ["leaf-1", "Healthy"], ["leaf-2", "Down"], ["spine-1", "Healthy"]),
    "chart_os_version": "Here is the chart. [Image saved: image_1.png]",
}
BAD = {
    "devices_list": T(["Hostname"], ["leaf-1"], ["leaf-2"], ["ghost-9"]),
    "devices_count": "You have 4 devices.",
    "devices_fields": T(["Hostname", "Mgmt IP", "Model", "Serial", "Version"],
                        ["leaf-1", "10.0.0.1", "N9K-C93180YC-FX", "WRONG", "10.3(3)"],
                        ["leaf-2", "10.0.0.2", "N9K-C93180YC-FX", "FDO222", "10.3(3)"],
                        ["spine-1", "10.0.0.3", "N9K-C9332C", "FDO333", "10.2(5)"]),
    "os_version_counts": T(["OS version", "Devices"], ["10.3(3)", 3]),
    "models_list": "Model in use: N9K-C9500.",
    "unhealthy_devices": "All devices look healthy.",
    "cpu_all": T(["Device", "CPU %"], ["leaf-1", 13], ["leaf-2", 50], ["spine-1", 44]),
    "mem_all": "Memory data is not available for this connector.",
    "cpu_mem_device": "Device leaf-1: CPU utilization 60%, memory utilization 41%.",
    "cpu_top": "leaf-1 has the highest CPU.",
    "cpu_above": T(["Device", "CPU"], ["leaf-2", "91%"], ["leaf-1", "85%"]),       # leaf-1 is at 12 %
    "mem_above": "No devices are above 75%.",
    "interfaces_list": "leaf-1 has no interfaces.",
    "interfaces_down": "No interfaces are down on leaf-1.",
    "interface_counters": "Interface counters are not available for leaf-1.",
    "links": T(["Local", "Remote"], ["leaf-1 Eth1/49", "leaf-2 Eth1/1"]),
    "fan_psu": T(["Device", "Part", "Status"], ["leaf-1", "Fan1", "OK"], ["leaf-2", "PSU2", "OK"], ["spine-1", "Fan1", "OK"]),
    "temperature": T(["Device", "Temperature (°C)"], ["leaf-1", 36], ["leaf-2", 70], ["spine-1", 41]),
    "health_summary": T(["Device", "Health"], ["leaf-1", "Healthy"], ["leaf-2", "Healthy"], ["spine-1", "Healthy"]),
    "chart_os_version": "10.3(3): 2 devices, 10.2(5): 1 device.",
}
# Part of the data, all of it right: PASS ("partial: ...") with PARTIAL_PASS on (Dev 2026-10-07),
# FAIL with it off (the old rule). Wrong / invented data FAILs either way (BAD above).
PARTIAL = {
    "devices_list": T(["Hostname", "IP"], ["leaf-1", "10.0.0.1"], ["leaf-2", "10.0.0.2"]),
    "devices_fields": T(["Hostname", "Mgmt IP", "Version"], ["leaf-1", "10.0.0.1", "10.3(3)"],
                        ["leaf-2", "10.0.0.2", "10.3(3)"]),                  # 2 of 3 devices, no model / serial
    "models_list": "Model in use: N9K-C93180YC-FX.",
    "cpu_all": T(["Device", "CPU %"], ["leaf-1", 13], ["leaf-2", 89.5]),
    "interfaces_list": "leaf-1 interfaces: Ethernet1/1.",
    "interface_counters": "Counters for Ethernet1/1: 100 bytes in.",
    "links": T(["Local", "Remote"], ["leaf-1 Eth1/49", "spine-1 Eth1/1"]),
    "fan_psu": T(["Device", "Part", "Status"], ["leaf-1", "Fan1", "OK"], ["leaf-2", "PSU2", "Failed"]),
    "health_summary": T(["Device", "Health"], ["leaf-1", "Healthy"], ["leaf-2", "Down"]),
}
PARAM = {"cpu_above": 80, "mem_above": 75}


def run(check, answer, src=None, has_image=None):
    src = src or FakeSource()
    row = PromptRow("PX", "prompt", check, PARAM.get(check))
    dev = src.devices()[0]
    return evaluate(Ctx(row, src, answer, has_image if has_image is not None else "[Image" in answer, dev))


def test_every_check_has_a_good_and_bad_sample():
    assert set(GOOD) == set(CHECKS) == set(BAD)


@pytest.mark.parametrize("check", sorted(GOOD))
def test_good_answer_passes(check):
    v = run(check, GOOD[check])
    assert v.status == "PASS", (v.reason, v.expected)


@pytest.mark.parametrize("check", sorted(BAD))
def test_bad_answer_fails(check):
    v = run(check, BAD[check])
    assert v.status == "FAIL", (v.reason, v.expected)


@pytest.mark.parametrize("check", sorted(PARTIAL))
def test_partial_answer_passes_only_with_partial_pass(check, monkeypatch):
    monkeypatch.setattr(settings, "PARTIAL_PASS", True)
    v = run(check, PARTIAL[check])
    assert v.status == "PASS" and v.reason.startswith("partial:"), (v.reason, v.expected)
    monkeypatch.setattr(settings, "PARTIAL_PASS", False)
    assert run(check, PARTIAL[check]).status == "FAIL"


def test_partial_pass_never_hides_wrong_or_invented_data(monkeypatch):
    monkeypatch.setattr(settings, "PARTIAL_PASS", True)
    assert run("devices_list", T(["Hostname"], ["leaf-1"], ["ghost-9"])).status == "FAIL"           # invented
    assert run("cpu_all", T(["Device", "CPU %"], ["leaf-1", 13], ["leaf-2", 50])).status == "FAIL"   # wrong value
    fields = T(["Hostname", "Model"], ["leaf-1", "N9K-C93180YC-FX"], ["leaf-2", "Cisco Nexus 9300 Switch"])
    assert run("devices_fields", fields).status == "FAIL"           # a model column with a wrong value


def test_answer_cut_off_by_the_timeout_is_graded():
    row = PromptRow("PX", "p", "devices_list")
    v = evaluate(Ctx(row, FakeSource(), GOOD["devices_list"]), "timed out after 300s")
    assert v.status == "PASS" and "graded the text that had arrived" in v.reason
    assert evaluate(Ctx(row, FakeSource(), ""), "timed out after 300s").status == "FAIL"


def test_single_device_values_skip_timestamps_and_prefer_percent():
    # zabbix-P09 (conv 333): "05:39" in the heading was read as CPU 5 / memory 5
    answer = ("**leaf-1 – CPU & Memory (as of 2026-10-07 05:39 UTC)**\n\n"
              + T(["Category", "Metric", "Value"], ["**CPU**", "CPU utilization", "**12.4 %**"],
                  ["**Memory**", "Total memory", "**4 080 189 440 B**"],
                  ["**Memory**", "Memory utilization", "**40.2 %**"]))
    v = run("cpu_mem_device", answer)
    assert v.status == "PASS", v.reason


def test_not_available_wording_with_curly_apostrophe_is_na():
    # ones-P12 (conv 367)
    answer = ("I wasn’t able to retrieve any memory‑utilization data from ONES‑MCP. The telemetry source "
              "does not currently expose memory‑usage metrics.")
    assert run("mem_all", answer, FakeSource(no_mem=True)).status == "NA"


def test_metric_matches_any_sample_taken_while_ncp_answered():
    class Moving(FakeSource):
        def __init__(self):
            super().__init__()
            self.reads = 0

        def _metrics(self):
            self.reads += 1
            cpu = 30 if self.reads == 1 else 70             # the value moved during the answer
            return {"leaf-1": {"cpu": cpu, "mem": 40, "temp": 35, "temp_alt": [35, 48]},
                    "leaf-2": {"cpu": 91, "mem": 80, "temp": 52}, "spine-1": {"cpu": 45, "mem": 60, "temp": 41}}
    src = Moving()
    src.begin_window(every=0.05)
    time.sleep(0.2)
    src.end_window()
    assert run("cpu_all", T(["Device", "CPU %"], ["leaf-1", 31], ["leaf-2", 91], ["spine-1", 45]), src).status == "PASS"
    assert run("cpu_all", T(["Device", "CPU %"], ["leaf-1", 50], ["leaf-2", 91], ["spine-1", 45]), src).status == "FAIL"
    # every temperature sensor counts (zabbix-P18: NCP gave another sensor than the hottest)
    assert run("temperature", T(["Device", "Temperature °C"], ["leaf-1", 48], ["leaf-2", 52], ["spine-1", 41]),
               src).status == "PASS"


def test_device_pick_skips_duplicate_names():
    class Dupes(FakeSource):
        def _devices(self):
            return [Device("Leaf-1", "10.4.4.64"), Device("Leaf-1", "10.4.6.11"), Device("Leaf-2", "10.4.6.12")]
    assert Dupes().device_for_prompts().name == "Leaf-2"


def test_zabbix_part_status_and_memory_pool():
    from ncp_suite.truth.zabbix import memory_value, part_ok
    assert [part_ok(t) for t in ("up", "on", "normal", "enabled", "down", "offEnvPower", "notPresent", "disabled")] \
        == [True, True, True, True, False, False, None, None]
    pools = [{"key_": "vm.memory.util[vm.memory.util.11]", "name": "reserve Processor: Memory utilization",
              "lastvalue": "0.086"},
             {"key_": "vm.memory.util[vm.memory.util.7]", "name": "IOS Process stack: Memory utilization",
              "lastvalue": "65.9"},
             {"key_": "vm.memory.util[vm.memory.util.1]", "name": "Processor: Memory utilization", "lastvalue": "31.8"}]
    assert memory_value(pools) == 31.8
    assert memory_value([{"key_": "sonic.snmp.mem.util", "name": "Memory utilization", "lastvalue": "47.2"}]) == 47.2
    # NCP's own words for a powered-off PSU count as "reported faulty" (zabbix-P17, conv 359)
    from ncp_suite.grading.compare import has_bad_word
    assert has_bad_word("| cisconx-engai-leaf01 | up (2) | offEnvPower (5) |")


def test_placeholder_row_is_not_an_invented_device(monkeypatch):
    # ones-P01 (conv 383): NCP ends a cut-short table with a filler row
    monkeypatch.setattr(settings, "PARTIAL_PASS", True)
    answer = T(["Hostname", "IP"], ["leaf-1", "10.0.0.1"], ["leaf-2", "10.0.0.2"],
               ["… *(additional rows omitted for brevity)*", ""])
    v = run("devices_list", answer)
    assert v.status == "PASS" and v.reason.startswith("partial:"), v.reason


def test_values_in_a_numbered_list_without_the_metric_word():
    # zabbix-P08 (conv 408): "1. Arista Leaf 1 – 96.15 %", plus a summary line naming two devices
    answer = ("**Memory utilization**\n- Range: **40.5 %** (lowest, *leaf-1*) → **80 %** (highest, *leaf-2*)\n"
              "1. leaf-2 – 80 %\n2. spine-1 – 60.4 %\n3. leaf-1 – 40.5 %")
    assert run("mem_all", answer).status == "PASS"


def test_a_column_ncp_did_not_show_is_not_shown_even_if_a_name_holds_an_ip(monkeypatch):
    # zabbix-P03 (conv 391): no IP column; "Linux Server 10.0.0.9" holds an IP in its name
    monkeypatch.setattr(settings, "PARTIAL_PASS", True)

    class Hosts(FakeSource):
        def _devices(self):
            return [Device("Linux Server 10.0.0.9", "10.0.0.9", "X1", "S1", "1.0"),
                    Device("leaf-1", "10.0.0.1", "N9K", "FDO1", "10.3")]
    answer = T(["Host", "model", "serial", "osversion"], ["Linux Server 10.0.0.9", "X1", "S1", "1.0"],
               ["leaf-1", "N9K", "FDO1", "10.3"])
    v = run("devices_fields", answer, Hosts())
    assert v.status == "PASS" and "field mgmt IP" in v.reason, v.reason


def test_models_fall_back_to_platform_when_the_source_has_no_model():
    class NoModel(FakeSource):          # ONES 10.20.0.37: model is random per read, so left out
        def _devices(self):
            return [Device("leaf-1", "10.0.0.1", "", "", "4.4", platform="wistron_6512_32r"),
                    Device("leaf-2", "10.0.0.2", "", "", "4.4", platform="Accton-AS7326-56X")]
    assert run("models_list", "Platforms: wistron_6512_32r and Accton-AS7326-56X.", NoModel()).status == "PASS"
    assert run("models_list", "Models: FOOS9443.", NoModel()).status == "FAIL"


def test_faulty_part_reported_with_the_sources_own_status_word():
    class OnesPsu(FakeSource):          # ones-P17 (conv 464): ONES PSU status "False"
        def _components(self):
            return [Component("leaf-1", "psu", "PSU 1", "True", True), Component("leaf-2", "psu", "PSU 1", "False", False)]
    answer = T(["Hostname", "PSU Name", "Operational Status"], ["leaf-1", "PSU 1", "True"], ["leaf-2", "PSU 1", "False"])
    assert run("fan_psu", answer, OnesPsu()).status == "PASS"


class MovingInventory(FakeSource):
    """ONES 10.20.0.37: every 60 s a new serial for most devices and new health for many."""
    def __init__(self):
        super().__init__()
        self.reads = 0

    def _devices(self):
        self.reads += 1
        first = self.reads == 1
        return [Device("leaf-1", "10.0.0.1", "N9K", "OLD1" if first else "NEW1", "10.3", healthy=not first),
                Device("leaf-2", "10.0.0.2", "N9K", "S2", "10.3", healthy=False, reason="PSU out of power"),
                Device("spine-1", "10.0.0.3", "N9K", "S3", "10.3", healthy=True)]


def _sampled(src):
    src.begin_window(every=0.05, devices=True)
    time.sleep(0.2)
    src.end_window()
    return src


def test_a_field_from_any_inventory_sample_during_the_answer_counts():
    src = _sampled(MovingInventory())              # leaf-1 serial: OLD1 first, NEW1 later
    answer = T(["Hostname", "Mgmt IP", "Model", "Serial", "Version"], ["leaf-1", "10.0.0.1", "N9K", "OLD1", "10.3"],
               ["leaf-2", "10.0.0.2", "N9K", "S2", "10.3"], ["spine-1", "10.0.0.3", "N9K", "S3", "10.3"])
    assert run("devices_fields", answer, src).status == "PASS"
    assert run("devices_fields", answer.replace("OLD1", "XXX9"), src).status == "FAIL"


def test_health_that_changed_during_the_answer_is_not_required():
    src = _sampled(MovingInventory())              # leaf-1 unhealthy at the first sample only
    v = run("unhealthy_devices", "leaf-2 is unhealthy: PSU out of power.", src)
    assert v.status == "PASS" and "changed health during the answer" in v.expected, (v.reason, v.expected)
    assert run("unhealthy_devices", "All devices look healthy.", src).status == "FAIL"    # leaf-2 stayed bad


def test_devices_listed_under_an_unhealthy_heading_are_flagged():
    # ones-P19 (conv 545): "### Unhealthy devices" then a comma-separated list of names
    answer = ("**Health summary**\n\n| Metric | Value |\n|---|---|\n| Unhealthy | 1 |\n\n"
              "### Unhealthy devices (reported by the controller)\n\nleaf-2,\n\n### Healthy devices\n\nleaf-1, spine-1")
    assert run("health_summary", answer).status == "PASS"
    assert run("unhealthy_devices", answer).status == "PASS"


def test_devices_shown_below_the_threshold_as_context_are_not_listed():
    # catalyst-P11 (conv 672): "none above 80 %" plus every device with its lower value (Dev, §11 q8)
    answer = "None of the devices is above 95 %.\n\n" + T(["Device", "CPU %"], ["leaf-1", 12], ["leaf-2", 91],
                                                         ["spine-1", 45])
    assert evaluate(Ctx(PromptRow("PX", "p", "cpu_above", 95), FakeSource(), answer)).status == "PASS"


def test_a_text_bar_chart_is_a_chart():
    # catalyst-P20 (conv 708): a bar chart drawn with block characters
    assert run("chart_os_version", "```\n10.3(3)  ████ (2)\n10.2(5)  ██ (1)\n```").status == "PASS"
    assert run("chart_os_version", "```\n10.3(3)  ████ (3)\n10.2(5)  ██ (1)\n```").status == "FAIL"   # wrong count


def test_chart_data_from_the_tool_call_is_checked():
    # zabbix-P20 (run 160637): generate_column_chart's arguments hold the plotted data
    def chart(*rows):
        return [{"tool_name": "UI_Visualization", "arguments": "{}"},
                {"tool_name": "generate_column_chart",
                 "arguments": json.dumps({"data": [{"category": c, "value": v} for c, v in rows]})}]
    row = PromptRow("P20", "p", "chart_os_version")
    text = "Here is the chart."                                   # no ui:// tag in the text: the trace shows it
    assert evaluate(Ctx(row, FakeSource(), text, trace=chart(("10.3(3)", 2), ("10.2(5)", 1)))).status == "PASS"
    assert evaluate(Ctx(row, FakeSource(), text, trace=chart(("10.3(3)", 3), ("10.2(5)", 1)))).status == "FAIL"
    assert evaluate(Ctx(row, FakeSource(), text)).status == "FAIL"                           # no chart at all


def test_a_problem_count_column_flags_the_device():
    # zabbix-P19 (conv 704): "Active problems" column instead of a health word
    table = lambda n: T(["Host", "CPU", "Active problems*"], ["leaf-1", "12 %", 0], ["leaf-2", "91 %", n],
                        ["spine-1", "45 %", 0])
    assert run("health_summary", table(2)).status == "PASS"
    assert run("health_summary", table(0)).status == "FAIL"


def test_a_platform_column_is_not_a_model_column(monkeypatch):
    # zabbix-P03 (conv 504): "Platform" holds EOS / IOS — NCP gave no model at all
    monkeypatch.setattr(settings, "PARTIAL_PASS", True)
    answer = T(["Host", "IP", "Vendor", "Platform"], ["leaf-1", "10.0.0.1", "Cisco", "NX-OS"],
               ["leaf-2", "10.0.0.2", "Cisco", "NX-OS"], ["spine-1", "10.0.0.3", "Cisco", "NX-OS"])
    v = run("devices_fields", answer)
    assert v.status == "PASS" and "field model" in v.reason, v.reason


def test_not_available_while_source_has_data_says_so():
    v = run("mem_all", BAD["mem_all"])
    assert v.status == "FAIL" and "not available" in v.reason


def test_unsupported_and_ncp_says_so_is_na():
    v = run("links", "Link / topology data is not available from Zabbix.", FakeSource(links_unsupported=True))
    assert v.status == "NA", v.reason


def test_unsupported_but_ncp_invents_data_fails():
    v = run("links", GOOD["links"], FakeSource(links_unsupported=True))
    assert v.status == "FAIL" and "invented" in v.reason


def test_no_ground_truth_is_blocked_not_fail():
    v = run("cpu_all", GOOD["cpu_all"], FakeSource(metrics_broken=True))
    assert v.status == "BLOCKED" and "404" in v.reason


def test_ncp_error_fails():
    row = PromptRow("PX", "p", "devices_list")
    assert evaluate(Ctx(row, FakeSource(), ""), "timed out after 120s").status == "FAIL"


def test_prompt_sheet_matches_checks():
    rows = load_prompts(settings.PROMPTS_XLSX)
    assert len(rows) == 20 and all(r.check in CHECKS for r in rows)


def test_followup_detection():
    assert is_followup("Which device would you like me to check?")
    assert not is_followup(GOOD["cpu_all"] + "\nWould you like a chart?")
    assert is_followup("The request to the connector failed. Shall I try the request again?")      # Ticketing
    # DC-INVENTORY: a definite "no data" answer is final; Ticketing: a chart is an answer
    assert not is_followup("I wasn't able to retrieve CPU data from ONES. Would you like another source?")
    assert not is_followup("Here is the chart:\n\n![](ui://bar-chart-1)\n\nWould you like it as a table?")


def test_followup_replies_start_with_the_tag():
    ctx = {"connector": "ONES", "tag": "#ONES-MCP", "device": "Leaf-1"}
    assert followup_reply("Which data source should I use?", ctx) == "#ONES-MCP Use the ONES connector for this."
    assert followup_reply("Which interface on Leaf-1 would you like?", ctx) == "#ONES-MCP All interfaces on Leaf-1."
    assert followup_reply("Which device would you like?", ctx) == "#ONES-MCP Device Leaf-1."
    assert followup_reply("Shall I go on?", ctx).startswith("#ONES-MCP Yes, please go ahead")


def test_which_one_question_with_a_table_is_a_followup_answered_with_the_ip():
    # ones-P13 (conv 370): two devices named Leaf-1, NCP asks which one, with a small table
    q = ("I found two devices named **Leaf‑1** in the ONES inventory:\n\n"
         + T(["Hostname", "Management IP"], ["Leaf-1", "10.4.4.64"], ["Leaf-1", "10.4.6.11"])
         + "\n\nWhich one would you like the interface list for? Please specify the IP address.")
    assert is_followup(q)
    ctx = {"connector": "ONES", "tag": "#ones-37-mcp", "device": "Leaf-1", "device_ip": "10.4.4.64"}
    assert followup_reply(q, ctx) == "#ones-37-mcp Device Leaf-1 (management IP 10.4.4.64)."
    # a full answer that ends with an offer is still an answer
    assert not is_followup(GOOD["cpu_all"] + "\n\nWould you like a chart of these values?")
    # ones-P14 (conv 460): NCP could not find the device by name and asks for its IP
    ask_ip = ("I’m unable to locate a device named **AS7326‑6045** in the ONES inventory. Could you provide the "
              "management IP address (the `switchip` value) for that switch? Once I have the IP, I can retrieve "
              "and show you all interfaces where `oper_status = down`.")
    assert is_followup(ask_ip)
    assert followup_reply(ask_ip, {**ctx, "device": "AS7326-6045", "device_ip": "10.20.12.32"}) \
        == "#ones-37-mcp Device AS7326-6045 (management IP 10.20.12.32)."
    # ones-P08 (conv 529): NCP asks to narrow a whole-fleet request; a user says "all of them"
    narrow = ("I can retrieve memory‑utilization values, but the ONES‑MCP connector only returns that metric on a "
              "per‑device basis. To avoid making hundreds of individual calls, could you narrow the request—e.g., "
              "give a list of device IPs/hostnames?")
    assert is_followup(narrow)
    assert followup_reply(narrow, {"connector": "ONES", "tag": "#ones-37-mcp"}) \
        == "#ones-37-mcp Yes, all devices in ONES, please. One call per device is fine."
    assert not is_followup(GOOD["cpu_all"] + "\n\nWould you like to narrow down the list by region?")   # an answer
    per_call = ("I can pull the current memory-utilization metric (`memUtil`) from each device with the ONES "
                "`get_device_system` call, but that requires a separate request per device. Shall I go ahead?")
    assert followup_reply(per_call, {"connector": "ONES", "tag": "#ones-37-mcp"}).endswith("One call per device is fine.")


# ---------------------------------------------------------------- chat client vs a fake NCP
def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


# NCP's table widget: the answer text only holds ![](ui://...), the rows come in a tool-result frame
WIDGET = {"type": "resource", "resource": {"uri": "ui://data-table-1", "mimeType": "text/html;profile=mcp-app",
                                           "meta": {"ui": {"templateUri": "ui://template/data-table"}}},
          "structuredContent": {"title": "CPU", "columns": ["Device", "CPU %"],
                                "rows": [["leaf-1", 13], ["leaf-2", 89.5], ["spine-1", 44]]}}


# agent_trace as NCP sends it on agent_complete (10.4.5.10, conversation 3243 shape); a switch login in a result
TRACE = [{"type": "tool_call", "agent_path": ["orchestrator"], "depth": 0, "tool_name": "query_fake",
          "connector_name": None, "arguments": '{"query": "How many devices?"}', "success": True,
          "result": "There are 3 devices.", "error": None, "duration_ms": 900},
         {"type": "tool_call", "agent_path": ["orchestrator", "fake_agent"], "depth": 1, "tool_name": "get_devices",
          "connector_name": "fake", "arguments": "{}", "success": True,
          "result": '{"devices": [{"hostname": "leaf-1", "username": "admin", "password": "S3cret!"}]}',
          "error": None, "duration_ms": 120},
         {"type": "tool_call", "agent_path": ["orchestrator", "fake_agent"], "depth": 1, "tool_name": "get_fabrics",
          "connector_name": "fake", "arguments": "{}", "success": True,
          "result": "API request failed: Server error '500 Internal Server Error' for url 'https://x/api/fabrics'",
          "error": None, "duration_ms": 80}]


async def _notifications(ws, seconds: float) -> None:
    """NCP sends notifications (no conversation_id) to every open socket of the user."""
    end = time.monotonic() + seconds
    try:
        while time.monotonic() < end:
            await asyncio.sleep(0.3)
            await ws.send(json.dumps({"type": "new_notification", "data": {"title": "x"}}))
    except websockets.exceptions.ConnectionClosed:      # the client is done and has hung up
        pass


async def _fake_ncp(ws):
    await ws.send(json.dumps({"type": "connection_id", "client_id": "c1"}))
    assert json.loads(await ws.recv())["type"] == "auth"
    await ws.send(json.dumps({"type": "auth_success"}))
    await ws.send(json.dumps({"type": "conversations_loaded"}))
    async for raw in ws:
        msg = json.loads(raw)
        if msg["type"] == "new_conversation":
            await ws.send(json.dumps({"type": "conversation_created", "conversation": {"id": 42}}))
        elif msg["type"] == "load_messages":
            await ws.send(json.dumps({"type": "messages_loaded", "messages": [{"role": "ASSISTANT",
                                                                               "ui_resources": [WIDGET]}]}))
        elif msg["type"] == "new_message":
            if msg["message"].startswith("#complete"):    # 10.4.5.236: end frame is agent_complete, then noise
                await ws.send(json.dumps({"type": "agent_tool_call", "conversation_id": 42,
                                          "tool_name": "get_devices"}))
                await ws.send(json.dumps({"type": "agent_llm_stream", "conversation_id": 42,
                                          "chunk": "There are 3 devices in your network."}))
                await ws.send(json.dumps({"type": "agent_complete", "conversation_id": 42, "agent_trace": TRACE}))
                await ws.send(json.dumps({"type": "follow_up_suggestions", "conversation_id": 42,
                                          "suggestions": ["Show CPU"]}))
                await _notifications(ws, 6)
                continue
            if msg["message"].startswith("#noise"):       # no end frame; another chat's frame; notifications
                await ws.send(json.dumps({"type": "agent_llm_stream", "conversation_id": 42,
                                          "chunk": "There are 3 devices in your network."}))
                await ws.send(json.dumps({"type": "agent_llm_stream", "conversation_id": 99, "chunk": " GHOST"}))
                await _notifications(ws, 6)
                continue
            if msg["message"].startswith("#widget"):      # "#widget saved": rows only in the saved message
                if "saved" not in msg["message"]:
                    await ws.send(json.dumps({"type": "agent_tool_result", "ui_resources": [WIDGET]}))
                await ws.send(json.dumps({"type": "agent_llm_stream", "chunk": "CPU table:\n\n![](ui://data-table-1)"}))
            elif msg["message"].startswith("#silent"):    # NCP never answers
                continue
            elif msg["message"].startswith("#trickle"):   # NCP is still writing when the timeout hits
                try:
                    for i in range(15):
                        await ws.send(json.dumps({"type": "agent_llm_stream", "conversation_id": 42,
                                                  "chunk": f"row {i}\n"}))
                        await asyncio.sleep(0.3)
                except websockets.exceptions.ConnectionClosed:
                    pass
                continue
            elif msg["message"].startswith("#fake") and "Device" not in msg["message"]:
                for chunk in ("Which device ", "would you like me to check?"):
                    await ws.send(json.dumps({"type": "agent_llm_stream", "chunk": chunk}))
            else:   # the follow-up reply
                table = GOOD["cpu_all"]
                for line in table.splitlines(keepends=True):
                    await ws.send(json.dumps({"type": "agent_llm_stream", "chunk": line}))
                await ws.send(json.dumps({"type": "new_message", "message": {
                    "contents": [{"content_type": "TEXT", "content": table}]}}))
            await ws.send(json.dumps({"type": "agent_completed"}))


@pytest.fixture
def fake_ncp(tmp_path):
    port, loop, stop = _free_port(), asyncio.new_event_loop(), None

    def serve():
        nonlocal stop
        asyncio.set_event_loop(loop)
        stop = loop.create_future()

        async def main():
            async with websockets.serve(_fake_ncp, "127.0.0.1", port):
                await stop
        loop.run_until_complete(main())

    thread = threading.Thread(target=serve, daemon=True)
    thread.start()
    time.sleep(0.5)
    yield ChatSettings(ws_uri=f"ws://127.0.0.1:{port}", login_url="", user="superadmin",
                       quiet_seconds=2, end_grace_seconds=0.5, retries=1, image_dir=tmp_path)
    loop.call_soon_threadsafe(stop.set_result, None)
    thread.join(timeout=5)


def _chat(cfg) -> NcpChat:
    chat = NcpChat(cfg)
    chat._token = "tok"
    return chat


def test_chat_with_followup(fake_ncp):
    res = _chat(fake_ncp).ask("#fake Show CPU utilization of all devices.",
                              {"connector": "Fake", "tag": "#fake", "device": "leaf-1"}, timeout=20)
    assert res.ok, res.error
    assert res.conversation_id == "42"
    assert len(res.turns) == 2 and res.turns[1][0] == "#fake Device leaf-1."   # the tag first (CHANGED 11)
    assert res.text.count("| leaf-2 |") == 1            # streamed + final text not doubled
    assert run("cpu_all", res.text).status == "PASS"


@pytest.mark.parametrize("prompt", ["#widget Show CPU.", "#widget saved Show CPU."])
def test_chat_table_widget_is_read(fake_ncp, prompt):
    res = _chat(fake_ncp).ask(prompt, {}, timeout=20)
    assert res.ok, res.error
    assert "ui://" not in res.text and "| leaf-2 | 89.5 |" in res.text
    assert run("cpu_all", res.text).status == "PASS"


def test_answer_ends_on_agent_complete_not_on_silence(fake_ncp):
    # quiet timer set high: only the end frame (+0.5 s grace) can end this answer quickly
    res = _chat(replace(fake_ncp, quiet_seconds=30)).ask("#complete How many devices?", {}, timeout=60)
    assert res.ok, res.error
    assert res.text == "There are 3 devices in your network."
    assert res.tools == ["get_devices"]
    assert res.seconds < 4, f"took {res.seconds}s: notifications after agent_complete kept the answer open"
    # the tool payload is kept, and the switch login in it is masked (rule 11)
    from ncp_suite.chat import trace as tr
    calls = tr.calls(res.trace)
    assert len(calls) == 3 and "S3cret" not in json.dumps(calls) and '"password": "***"' in calls[1]["result"]
    assert tr.summary(calls).startswith("2 data call(s) · 1 failed: get_fabrics → ")


def test_notifications_and_other_chats_do_not_keep_answer_open_or_leak(fake_ncp):
    # no end frame: silence after text (2 s) ends it, although notifications keep coming for 6 s
    res = _chat(fake_ncp).ask("#noise How many devices?", {}, timeout=60)
    assert res.ok, res.error
    assert "GHOST" not in res.text and res.text == "There are 3 devices in your network."
    assert res.seconds < 5, f"took {res.seconds}s: notifications counted as activity"


def test_an_answer_timeout_is_repeated_once_and_listed(fake_ncp):
    res = _chat(replace(fake_ncp, retries=3, answer_retries=1)).ask("#silent Show CPU.", {}, timeout=1)
    assert error_kind(res.error) == "answer" and len(res.retries) == 1, (res.error, res.retries)
    assert res.retries[0].startswith("attempt 1: timed out after 1s")


def test_an_answer_cut_off_with_text_is_kept_not_repeated(fake_ncp):
    res = _chat(replace(fake_ncp, retries=3, answer_retries=1)).ask("#trickle Show CPU.", {}, timeout=1)
    assert res.error == "timed out after 1s" and "row 0" in res.text and not res.retries, (res.error, res.retries)


def test_error_kinds():
    assert error_kind("timed out after 180s") == "answer"
    assert error_kind("no answer text received") == "answer"
    assert error_kind("ConnectionRefusedError: [Errno 61] Connect call failed") == "connection"
    assert error_kind("TimeoutError: timed out during opening handshake") == "connection"
    assert error_kind("NCP said something odd") is None


def test_answer_stream_rules():
    a = AnswerStream("42", lambda data, caption: "")
    assert not a.feed({"type": "new_notification"})                       # not activity
    assert not a.feed({"type": "agent_llm_stream", "conversation_id": 7, "chunk": "x"})   # other chat
    assert a.feed({"type": "agent_llm_stream", "conversation_id": "42", "chunk": "ok"})
    assert a.feed({"type": "agent_complete", "conversation_id": 42}) and a.ended
    assert a.feed({"type": "agent_started", "conversation_id": 42}) and not a.ended   # a later agent step
    assert a.text() == "ok"


def test_chart_widget_is_not_turned_into_a_table():
    chart = {"resource": {"uri": "ui://bar-chart-1", "meta": {"ui": {"templateUri": "ui://template/bar-chart"}}},
             "structuredContent": {"columns": ["OS", "Count"], "rows": [["9.3", 2]]}}
    text = "Chart:\n\n![](ui://bar-chart-1)"
    assert inline_ui(text, {"ui://bar-chart-1": chart}) == text


def test_non_breaking_hyphens_in_names_still_match():
    # NCP's LLM writes leaf‑1 (non-breaking hyphen) and 13 % (narrow no-break space)
    answer = GOOD["cpu_all"].replace("-", "‑").replace(" |", " |")
    assert run("cpu_all", answer).status == "PASS"


# ---------------------------------------------------------------- runner (no network)
class ScriptedChat:
    """Stands in for NcpChat: returns a fixed answer and remembers what was asked."""

    def __init__(self, text: str):
        self.text, self.asked = text, []

    def ask(self, prompt, context=None, timeout=None):
        from ncp_suite.chat import ChatResult
        self.asked.append((prompt, context, timeout))
        return ChatResult(prompt, text=self.text, seconds=1.5, conversation_id="77",
                          turns=[(prompt, self.text)], tools=["list_devices"])


FAKE_CONN = settings.Connector("fake", "Fake (Local MCP)", "#fake", "x:y", "http://x", "u", "p")


def test_runner_fills_device_sends_tag_and_grades():
    chat = ScriptedChat("Device leaf-1: CPU utilization 13%, memory utilization 41%.")
    row = PromptRow("P09", "Show CPU and memory for device <DEVICE>.", "cpu_mem_device", timeout=300)
    r = run_case(row, FAKE_CONN, chat, FakeSource())
    assert chat.asked[0][0] == "#fake Show CPU and memory for device leaf-1."
    assert chat.asked[0][2] == 300 and chat.asked[0][1]["device"] == "leaf-1"
    assert (r.status, r.device, r.conversation_id, r.tools) == ("PASS", "leaf-1", "77", ["list_devices"])
    assert "| leaf-1 | 12 | 40 |" in r.source                  # the graded source values, next to the answer


def test_source_view_shows_what_the_check_compared():
    from ncp_suite.grading.source_view import source_view
    src = FakeSource()
    assert "| leaf-2 | 10.0.0.2 | N9K-C93180YC-FX |" in source_view("devices_fields", src)
    assert "UNHEALTHY (down)" in source_view("unhealthy_devices", src)
    assert "| Ethernet1/2 | down |" in source_view("interfaces_down", src, src.devices()[0])
    assert "| leaf-2 | psu | PSU2 | failed | NO |" in source_view("fan_psu", src)
    assert source_view("links", FakeSource(links_unsupported=True)).endswith("has no link data.")
    assert source_view("cpu_all", FakeSource(metrics_broken=True)).startswith("Could not read the source")


def test_runner_known_issue_turns_fail_into_xfail():
    row = PromptRow("P02", "How many devices are in my network?", "devices_count", known_issue="NCP-123")
    assert run_case(row, FAKE_CONN, ScriptedChat("You have 4 devices."), FakeSource()).status == "XFAIL"


def test_runner_blocks_when_device_cannot_be_chosen():
    class NoDevices(FakeSource):
        def _devices(self):
            return []
    chat = ScriptedChat("unused")
    r = run_case(PromptRow("P13", "List all interfaces of device <DEVICE>.", "interfaces_list"),
                 FAKE_CONN, chat, NoDevices())
    assert r.status == "BLOCKED" and "<DEVICE>" in r.reason and not chat.asked


class SequenceChat(ScriptedChat):
    """Answers with the next text each time (a new chat per ask)."""
    def __init__(self, *texts):
        super().__init__(texts[0])
        self.texts = list(texts)

    def ask(self, prompt, context=None, timeout=None):
        self.text = self.texts[min(len(self.asked), len(self.texts) - 1)]
        res = super().ask(prompt, context, timeout)
        res.conversation_id = str(80 + len(self.asked))
        return res


def test_a_fail_is_repeated_once_and_a_pass_on_repeat_is_flaky(monkeypatch):
    monkeypatch.setattr(settings, "REPEAT_FAILS", 1)
    row = PromptRow("P02", "How many devices are in my network?", "devices_count")
    flaky = run_case(row, FAKE_CONN, SequenceChat("You have 4 devices.", "There are 3 devices."), FakeSource())
    assert flaky.status == "PASS" and flaky.reason.startswith("flaky: failed first (conversation 81")
    assert flaky.conversation_id == "82" and flaky.retries[0].startswith("attempt 1: FAIL")
    twice = run_case(row, FAKE_CONN, SequenceChat("You have 4 devices.", "You have 5 devices."), FakeSource())
    assert twice.status == "FAIL" and twice.reason.startswith("failed twice (conversations 81, 82)")
    monkeypatch.setattr(settings, "REPEAT_FAILS", 0)
    chat = SequenceChat("You have 4 devices.", "There are 3 devices.")
    assert run_case(row, FAKE_CONN, chat, FakeSource()).status == "FAIL" and len(chat.asked) == 1


def test_llm_reader_only_reads_and_the_code_still_grades(monkeypatch):
    from ncp_suite import runner
    monkeypatch.setattr(settings, "REPEAT_FAILS", 0)
    monkeypatch.setattr(runner, "extract_enabled", lambda check: True)
    row = PromptRow("P07", "Show CPU utilization of all devices.", "cpu_all")
    odd = "CPU looks fine: the first leaf sits near thirteen, leaf‑2 is busy and the spine is moderate."
    # the reader's table matches the source -> PASS, said so, table kept for the reviewer
    monkeypatch.setattr(runner, "extract_table", lambda c, q, a: T(["Device", "CPU %"], ["leaf-1", 13], ["leaf-2", 90],
                                                                      ["spine-1", 44]))
    r = run_case(row, FAKE_CONN, ScriptedChat(odd), FakeSource())
    assert r.status == "PASS" and "read via the LLM reader" in r.reason and "| leaf-1 | 13 |" in r.extracted
    # the reader's table is wrong -> the code still says FAIL
    monkeypatch.setattr(runner, "extract_table", lambda c, q, a: T(["Device", "CPU %"], ["leaf-1", 13], ["leaf-2", 50],
                                                                      ["spine-1", 44]))
    assert run_case(row, FAKE_CONN, ScriptedChat(odd), FakeSource()).status == "FAIL"


def test_llm_reader_reply_becomes_a_table():
    from ncp_suite.grading.extract import COLUMNS, _rows, to_table
    content = 'Sure:\n{"rows": [{"device": "leaf-1", "ip": null, "value": "13 %"}, {"device": "leaf-2", "value": 91}]}'
    assert to_table(_rows(content), COLUMNS["cpu_all"]) == ("| Device | IP | CPU % |\n|---|---|---|\n"
                                                            "| leaf-1 |  | 13 % |\n| leaf-2 |  | 91 |")
    assert to_table(_rows("no json here"), COLUMNS["cpu_all"]) == ""


# ---------------------------------------------------------------- source HTTP plumbing
def _response(status: int, body: bytes = b"{}"):
    import requests
    r = requests.Response()
    r.status_code, r._content = status, body
    return r


def test_source_logs_in_again_once_on_401():
    src, logins = FakeSource(), []
    src.login = lambda: logins.append(1)
    answers = iter([_response(401), _response(200, b'{"data": [1]}')])
    src.http.request = lambda *a, **kw: next(answers)
    assert src.get("/api/x") == {"data": [1]} and len(logins) == 2


def test_device_rows_without_a_name():
    class Nameless(FakeSource):
        def __init__(self, names):
            super().__init__()
            self.names = names

        def _devices(self):
            return [Device(n) for n in self.names]
    assert [d.name for d in Nameless(["leaf-1", ""]).devices()] == ["leaf-1"]
    with pytest.raises(NoTruth, match="has a name"):
        Nameless(["", " "]).devices()


# ---------------------------------------------------------------- settings and prompt sheet
def test_connectors_are_built_from_env_keys_by_name(monkeypatch):
    for key, value in {"TAG_DEMO": "#demo", "DEMO_URL": "https://demo", "DEMO_USER": "u",
                       "DEMO_PASSWORD": "p", "DEMO_SITE": "lab"}.items():
        monkeypatch.setenv(key, value)
    conn = settings._connector("demo", "Demo", "demo:DemoSource", {"site": ("DEMO_SITE", "")})
    assert (conn.tag, conn.url, conn.source, conn.extra) == ("#demo", "https://demo",
                                                             "ncp_suite.truth.demo:DemoSource", {"site": "lab"})
    assert [c.title for c in settings.CONNECTORS.values()] == ["Nexus Dashboard (Local MCP)",
                                                               "Catalyst Center (Local MCP)", "Zabbix", "ONES"]


def test_missing_names_blank_keys_never_values():
    conn = settings.Connector("demo", "Demo", "", "x:y", "https://demo", "user", "")
    blank = settings.missing([conn], ncp=False)
    assert blank == ["DEMO_PASSWORD"]
    assert "TAG_DEMO" in settings.missing([conn], ncp=True)


def test_prompt_sheet_timeout_column_is_optional(tmp_path):
    assert all(r.timeout is None for r in load_prompts(settings.PROMPTS_XLSX))
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "prompts"
    ws.append(["ID", "Prompt", "Check", "Timeout"])
    ws.append(["P01", "List all devices in my network.", "devices_list", 300])
    wb.save(tmp_path / "sheet.xlsx")
    assert load_prompts(tmp_path / "sheet.xlsx")[0].timeout == 300


# ---------------------------------------------------------------- report
def _row(pid, conn, status, reason="", answer="…", **kw):
    titles = {c.key: c.title for c in settings.CONNECTORS.values()}
    return PromptResult(id=pid, connector=conn, title=titles[conn], status=status, reason=reason,
                        expected="3 devices", sent=f"#tag prompt {pid}", answer=answer, seconds=1.0,
                        conversation_id="7", followups=kw.get("followups", []), tools=kw.get("tools", []),
                        retries=kw.get("retries", []))


def test_report_matrix_matches_devs_sheet(tmp_path):
    prompts = load_prompts(settings.PROMPTS_XLSX)[:2]
    conns = list(settings.CONNECTORS.values())
    results = [_row("P01", "ones", "PASS", "all 3 devices listed"),            # arrives first, sorted after
               _row("P01", "zabbix", "FAIL", "missing 1 of 3 devices: spine-1", tools=["host.get"],
                    answer="table \x1b[0m with a control character", retries=["attempt 1: timed out after 180s"])]
    path = write_report(results, prompts, conns, tmp_path)
    wb = load_workbook(path)
    ws = wb["Matrix"]
    assert [c.value for c in ws[1]] == ["Prompt", "Nexus Dashboard (Local MCP)", "Catalyst Center (Local MCP)",
                                        "Zabbix", "ONES", "Comments"]
    assert ws["D2"].value == "FAIL" and ws["E2"].value == "PASS" and "spine-1" in ws["F2"].value
    details = wb["Details"]
    assert [details["B2"].value, details["B3"].value] == ["Zabbix", "ONES"]   # matrix order, not arrival order
    assert details["L1"].value == "Tools called" and details["L2"].value == "host.get"
    assert details["M1"].value == "Retries" and "timed out" in details["M2"].value
    failures = wb["Failures"]                       # FAIL rows only, control character removed
    assert failures.max_row == 2 and failures["B2"].value == "Zabbix" and "\x1b" not in failures["L2"].value
    assert details["O1"].value == "Source data" and failures["K1"].value == "Source data"


def test_result_survives_the_trip_between_workers():
    r = _row("P20", "ones", "PASS", followups=[("a", "b")], tools=["t"])
    assert PromptResult.from_dict(json.loads(json.dumps(r.to_dict()))).to_dict() == json.loads(json.dumps(r.to_dict()))


# ---------------------------------------------------------------- HTML report parts
def test_html_matrix_has_devs_layout_counts_and_blanks():
    from ncp_suite.reporting.html import matrix_html
    prompts = load_prompts(settings.PROMPTS_XLSX)[:2]
    conns = list(settings.CONNECTORS.values())
    results = [_row("P01", "zabbix", "FAIL", "missing 1 of 3 devices: spine-1"), _row("P01", "ones", "PASS"),
               _row("P02", "nexus", "BLOCKED", "no ground truth"), _row("P02", "catalyst", "NA", "no such data")]
    page = matrix_html(results, prompts, conns, "x.xlsx")
    heads = ["Nexus Dashboard (Local MCP)", "Catalyst Center (Local MCP)", "Zabbix", "ONES", "Comments"]
    assert all(h in page for h in heads)
    assert page.index("Nexus Dashboard") < page.index("Catalyst Center") < page.index(">Zabbix<") < page.index(">ONES<")
    assert prompts[0].prompt in page and "spine-1" in page and "x.xlsx" in page
    # 8 planned cells, 4 results -> 4 "not run" across the connectors
    assert page.count('<td class="st"></td>') == 4


def test_reports_show_the_tool_trace_and_the_llm_reader_table(tmp_path):
    from ncp_suite.reporting.html import details_html
    calls = [{"path": "orchestrator > fake_agent", "tool": "get_fabrics", "connector": "fake", "arguments": "{}",
              "success": True, "error": "", "ms": 80, "result": "500 Internal Server Error <b>"}]
    row = _row("P07", "zabbix", "PASS", reason="ok — read via the LLM reader")
    row.trace, row.trace_summary = calls, "1 data call(s) · 1 failed: get_fabrics → 500"
    row.extracted = "| Device | CPU % |\n|---|---|\n| leaf-1 | 13 |"
    page = details_html(row, tmp_path)
    assert "1 failed: get_fabrics" in page and "&lt;b&gt;" in page
    assert "LLM reader" in page and "| leaf-1 | 13 |" in page
    wb = load_workbook(write_report([row], load_prompts(settings.PROMPTS_XLSX), list(settings.CONNECTORS.values()),
                                    tmp_path))
    head = [c.value for c in wb["Details"][1]]
    assert head[-2:] == ["NCP tool calls (agent_trace)", "LLM reader table"]
    assert "get_fabrics" in wb["Details"].cell(row=2, column=len(head) - 1).value


def test_html_details_escapes_answer_and_embeds_chart(tmp_path):
    from ncp_suite.reporting.html import details_html
    (tmp_path / "image_1.png").write_bytes(b"\x89PNG fake")
    row = _row("P20", "ones", "PASS", answer="<script>x</script> chart: [Image saved: image_1.png]",
               followups=[("Use the ONES connector (#ones) for this.", "Here it is")], tools=["render_chart"])
    page = details_html(row, tmp_path)
    assert "<script>x</script>" not in page and "&lt;script&gt;" in page
    assert "data:image/png;base64," in page and "Suite replied:" in page and "3 devices" in page
    assert "render_chart" in page and "Source (ground truth)" in page
