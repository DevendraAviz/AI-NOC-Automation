"""Offline self-tests for the GPU-metric connectors (Prometheus, DCGM; 2026-10-08). Fake sources, no network.
Every GPU check has an answer that must PASS (or NA, for data the lab does not export) and one that must FAIL."""
from __future__ import annotations

import pytest

from ncp_suite import settings
from ncp_suite.grading import gpu
from ncp_suite.grading.checks import CHECKS, Ctx, evaluate
from ncp_suite.prompts import PromptRow, load_prompts
from ncp_suite.truth.base import Unsupported
from ncp_suite.truth.prometheus import DcgmSource, Gpu, PrometheusSource

pytestmark = pytest.mark.offline


def T(header, *rows):
    return "\n".join(["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
                     + ["| " + " | ".join(map(str, r)) + " |" for r in rows])


GPUS = [Gpu("hgx-a", "0", "NVIDIA RTX A6000", "GPU-a0", "hgx-a", "10.0.0.1"),
        Gpu("hgx-a", "1", "NVIDIA RTX A6000", "GPU-a1", "hgx-a", "10.0.0.1"),
        Gpu("hgx-b", "0", "Tesla V100-SXM2-16GB", "GPU-b0", "10.0.0.9:9400", "10.0.0.9")]
NOW = {"hgx-a:0": {"util": 0, "temp": 40, "power": 100, "mem_pct": 10.0, "copy_util": 0, "xid": 0},
       "hgx-a:1": {"util": 95, "temp": 90, "power": 300, "mem_pct": 90.0, "copy_util": 40, "xid": 0},
       "hgx-b:0": {"util": 3, "temp": 50, "power": 60, "mem_pct": 0.0, "copy_util": 0}}
OVER = {("avg", "util"): [2, 80, 1], ("max", "util"): [5, 100, 3], ("min", "util"): [0, 60, 0],
        ("avg", "temp"): [40, 85, 50], ("max", "temp"): [45, 95, 52], ("min", "temp"): [35, 80, 48],
        ("avg", "power"): [100, 280, 50], ("max", "power"): [120, 310, 55], ("min", "power"): [90, 250, 45]}
NAMES = ["DCGM_FI_DEV_GPU_TEMP", "DCGM_FI_DEV_GPU_UTIL", "DCGM_FI_DEV_POWER_USAGE", "DCGM_FI_PROF_NVLINK_RX_BYTES",
         "node_cpu_seconds_total", "vllm:num_requests_running", "up"]


def _conn(key, title):
    return settings.Connector(key, title, f"#{key}", "x:y", "http://x", "", "", needs_login=False)


class FakeProm(PrometheusSource):
    def __init__(self, alerts=(), throttle=None):
        super().__init__(_conn("prometheus", "Prometheus (Local MCP)"))
        self._alerts, self._throttle = list(alerts), throttle

    def names(self):
        return NAMES + (["DCGM_FI_DEV_CLOCK_THROTTLE_REASONS"] if self._throttle else [])

    def gpus(self):
        return GPUS

    def _metrics(self, at=None):
        if at:                                     # the read at NCP's tool-call time: hgx-a GPU 1 was at 70 °C then
            return {k: dict(v, temp=70 if k == "hgx-a:1" else v["temp"]) for k, v in NOW.items()}
        return {k: dict(v) for k, v in NOW.items()}

    def over_time(self, fn, kind, hours):
        return {g.key: v for g, v in zip(GPUS, OVER.get((fn, kind), [])) if v is not None}

    def window_stats(self, kind, hours):
        return {g.key: {fn: [OVER[(fn, kind)][i]] for fn in ("avg", "max", "min")} for i, g in enumerate(GPUS)}

    def query(self, expr):
        return [({"Hostname": g.host, "gpu": g.index}, self._throttle.get(g.key, 0)) for g in GPUS]

    def alerts(self):
        return self._alerts


class FakeDcgm(FakeProm, DcgmSource):
    def __init__(self):
        super().__init__()
        self.title = "DCGM"

    alerts = DcgmSource.alerts


P, D = FakeProm, FakeDcgm
G = ["Host", "GPU"]
CASES = [
    ("metric_catalog", P, "Metrics: " + ", ".join(NAMES[:-1]) + ", `up`", "PASS"),
    ("metric_catalog", P, "Metrics: DCGM_FI_DEV_GPU_TEMP and DCGM_FI_DEV_GPU_MADE_UP_METRIC.", "FAIL"),
    # DCGM lists only its supported metrics; a real Prometheus metric outside that list is not "invented"
    ("metric_catalog", D, "DCGM exposes DCGM_FI_DEV_GPU_TEMP, DCGM_FI_DEV_GPU_UTIL, DCGM_FI_DEV_POWER_USAGE "
                          "and DCGM_FI_PROF_NVLINK_RX_BYTES.", "PASS"),
    ("gpu_metric_names", P, "GPU metrics: DCGM_FI_DEV_GPU_TEMP, DCGM_FI_DEV_GPU_UTIL, DCGM_FI_DEV_POWER_USAGE, "
                            "DCGM_FI_PROF_NVLINK_RX_BYTES.", "PASS"),
    ("gpu_metric_names", P, "No GPU metrics are available.", "FAIL"),
    ("gpu_util_chart", P, "Here is the chart. [Image saved: image_1.png]\nhgx-a averaged 41% utilization (peak 100%).", "PASS"),
    ("gpu_util_chart", P, "GPU utilization was low across the fleet.", "FAIL"),
    ("gpu_util_chart", P, "Chart: [Image saved: image_1.png]\nhgx-a averaged 5% utilization.", "FAIL"),
    ("power_temp_window", P, T(["Host", "Avg temp (°C)", "Max power (W)"], ["hgx-a", 62, 310], ["hgx-b", 50, 52]), "PASS"),
    ("power_temp_window", P, T(["Host", "Avg temp (°C)"], ["hgx-a", 20]), "FAIL"),
    # dcgm-G05 (conv 784): per GPU Avg / Min / Max / First / Last columns — each against its own statistic
    ("power_temp_window", P, "**GPU power usage – watts (W)**\n" + T(G + ["Avg W", "Min W", "Max W", "Last W"],
                                                                      ["hgx-a", 1, 281, 251, 309, 300]), "PASS"),
    ("power_temp_window", P, "**GPU power usage – watts (W)**\n" + T(G + ["Avg W", "Min W", "Max W", "Last W"],
                                                                      ["hgx-a", 1, 281, 120, 309, 300]), "FAIL"),
    ("power_temp_window", P, "**GPU power usage – watts (W)**\n" + T(G + ["Avg W", "Min W", "Max W"],
                                                                      ["hgx-a", 1, 281, 262, 296]), "PASS"),  # sampled min / max
    ("power_temp_window", P, "**GPU power usage – watts (W)**\n" + T(G + ["Avg W", "Last W"], ["hgx-a", 1, 281, 400]),
     "FAIL"),                                                                       # last outside the window's min–max
    ("power_temp_window", P, "Power over 24 h: hgx-a GPU 1 peaks at 310 W and hgx-b at 55 W, while idle GPUs stay "
                             "below 20 W. ![](ui://line-chart-1)", "PASS"),           # two hosts in one sentence
    # prometheus-G05 (conv 783): "X and Y … reach ~99 W and ~286 W respectively" — 286 W fits neither
    ("power_temp_window", P, "Power: hgx-b GPU 0 and hgx-a GPU 0 have the highest peaks, reaching ~55 W and ~500 W "
                             "respectively. ![](ui://line-chart-1)", "FAIL"),
    ("power_temp_window", P, "GPU power over 24 h. ![](ui://line-chart-1)\nThe hgx-b GPU 0 and hgx-a GPU 0 have the "
                             "highest peaks, reaching ~55 W and ~500 W respectively.", "FAIL"),   # no "power" word
    ("gpu_util_chart", P, "Average utilization chart: ![](ui://bar-chart-1)\nhgx-a ranges from 2 % on GPU 0 up to "
                          "80 % on GPU 1.", "PASS"),                                    # "up to" = spread, not a max
    ("gpu_temp_power", P, T(G + ["Temp (°C)", "Power (W)"], ["hgx-a", 0, 40, 100], ["hgx-a", 1, 90, 300],
                            ["hgx-b", 0, 50, 60]), "PASS"),
    ("gpu_temp_power", P, "hgx-a GPU 0 at 40 °C and GPU 1 at 90 °C; hgx-b at 50 °C.", "PASS"),
    ("gpu_temp_power", P, T(G + ["Temp (°C)", "Power (W)"], ["hgx-a", 0, 40, 100], ["hgx-a", 1, 60, 300]), "FAIL"),
    # prometheus-G06 (conv 469): temperatures in a 'Value' column under a 'Temperature and Power' title, and prose
    # that calls them power ("lowest power draw 28 W") -> the wrong power value FAILs
    ("gpu_temp_power", P, "Lowest power draw: **58 W** (`hgx-b`, GPU 0).\n\n**Current GPU Temperature and Power**\n\n"
     + T(["Hostname", "gpu", "Value"], ["hgx-a", 0, 40], ["hgx-a", 1, 90], ["hgx-b", 0, 50]), "PASS"),
    ("gpu_temp_power", P, "Lowest power draw: **28 W** (`hgx-b`, GPU 0).\n\n**Current GPU Temperature and Power**\n\n"
     + T(["Hostname", "gpu", "Value"], ["hgx-a", 0, 40], ["hgx-a", 1, 90], ["hgx-b", 0, 50]), "FAIL"),
    ("gpu_mem_util", P, T(G + ["Memory used %"], ["hgx-a", 0, "10%"], ["hgx-a", 1, "90%"], ["hgx-b", 0, "0%"]), "PASS"),
    ("gpu_mem_util", P, T(G + ["Memory used %"], ["hgx-a", 0, "10%"], ["hgx-a", 1, "50%"]), "FAIL"),
    ("gpu_mem_util", P, T(G + ["Memory copy util (bandwidth) %"], ["hgx-a", 1, "40%"]), "PASS"),
    ("throttling", P, "Throttling status is not available: no clock-throttle metric is exported.", "NA"),
    ("throttling", P, "No GPUs are currently throttled.", "FAIL"),
    ("throttling", P, "The throttle-reason field isn't included in this snapshot.", "NA"),
    # Vishakh 2026-10-08: no throttling series -> NCP's per-GPU temp / util / power are compared with the source
    ("throttling", P, "hgx-a GPU 1 at 92 °C with a reduced SM clock — a classic sign of thermal throttling. The "
                      "throttle-reason field isn't included in this snapshot.", "PASS"),
    ("throttling", P, T(G + ["Temp (°C)", "Util (%)", "Power (W)", "Likely throttling?"],
                        ["hgx-a", 1, 60, 95, 300, "Possible"]), "FAIL"),                      # wrong temperature
    ("ecc", D, "ECC error counters are not exported by this DCGM setup.", "NA"),
    ("ecc", D, "All GPUs show 0 single-bit and 0 double-bit ECC errors.", "FAIL"),
    ("gpu_hottest", P, "The hottest GPU is hgx-a GPU 1 at 90 °C, then hgx-b GPU 0 (50 °C).", "PASS"),
    ("gpu_hottest", P, "hgx-b GPU 0 is the hottest at 50 °C.", "FAIL"),
    ("gpu_hottest", P, "hgx-a is the hottest host.", "FAIL"),
    ("gpu_hottest", P, "The hottest is GPU 1 on hgx-a (90 °C).", "PASS"),
    ("gpu_hottest", P, "**Hottest GPU**\n" + T(["Host", "GPU #", "Temperature (°C)"], ["hgx-a", 1, "**90 °C**"]), "PASS"),
    ("gpu_hottest", P, "**Hottest GPU**\n" + T(["Host", "GPU #", "Temperature (°C)"], ["hgx-a", 0, "40 °C"]), "FAIL"),
    ("gpu_rank_util", P, T(["Rank"] + G + ["Util %"], [1, "hgx-a", 1, "95%"], [2, "hgx-b", 0, "3%"], [3, "hgx-a", 0, "0%"]),
     "PASS"),
    ("gpu_rank_util", P, T(G + ["Util %"], ["hgx-a", 0, "0%"], ["hgx-a", 1, "95%"], ["hgx-b", 0, "3%"]), "FAIL"),
    ("gpu_idle", P, "Idle GPUs: hgx-a GPU 0 (0%) and hgx-b GPU 0 (3%).", "PASS"),
    ("gpu_idle", P, T(G + ["Util %", "State"], ["hgx-a", 0, "0%", "idle"], ["hgx-a", 1, "95%", "busy"],
                      ["hgx-b", 0, "3%", "idle"]), "PASS"),                    # a busy row is context, not a claim
    ("gpu_idle", P, "Idle GPUs: hgx-a GPU 1 and hgx-b GPU 0.", "FAIL"),
    ("alerts_by_severity", P, "There are no firing alerts right now.", "PASS"),
    ("alerts_chart", P, "There are currently no active (firing) alerts in Prometheus, so a pie chart can't be "
                        "generated at this time.", "PASS"),                                   # conv 1127
    ("alerts_chart", P, "I wasn't able to find any time-series that match ALERTS{alertstate=\"active\"}, so there's "
                        "nothing to chart.", "PASS"),                                         # conv 1170
    # conv 1071: host-level values ("all GPUs on <host> at 100 %") are compared with every GPU of the host
    ("throttling", P, "All GPUs on the host hgx-a are at 50 % utilization, indicating they are throttled.", "FAIL"),
    ("throttling", P, "All GPUs on the host hgx-b are at 3 % utilization, meaning they are idle and not throttled.",
     "PASS"),
    ("alerts_by_severity", P, "Critical: GPUHighTemp on hgx-a.", "FAIL"),
    ("alerts_by_severity", D, "Alerts are not available from the DCGM connector.", "NA"),
    ("alerts_by_severity", D, "No alerts are firing.", "PASS"),            # Vishakh 2026-10-08: zero alerts = PASS
    ("alerts_by_severity", D, "Critical: GPUHot on hgx-a.", "FAIL"),       # DCGM has no alerts: an alert is invented
    ("alerts_chart", P, "No alerts are firing, so there is nothing to chart.", "PASS"),
    ("alerts_chart", P, "Here is the pie chart of alerts. [Image saved: image_2.png]", "FAIL"),
    # dcgm-G16 (conv 962): a pie chart with every severity 0, "no active alerts reported by DCGM" -> PASS
    ("alerts_chart", D, "Here's the pie chart: ![](ui://pie-chart-1)\nAll three severity levels—Critical, Warning, "
                        "and Info—currently have a count of 0, indicating that there are no active alerts.", "PASS"),
    ("alerts_chart", P, "![](ui://pie-chart-1) Critical, Warning and Info are all 0 — no active alerts.", "PASS"),
    ("alerts_chart", D, "DCGM does not provide alerts, so no chart can be drawn.", "NA"),
    ("gpu_fleet_health", P, "3 GPUs on 2 hosts.\n" + T(["Host", "GPUs", "Status"],
                                                       ["hgx-a", 2, "GPU 1 hot (90 °C) — warning"], ["hgx-b", 1, "healthy"]),
     "PASS"),
    ("gpu_fleet_health", P, T(["Host", "Status"], ["hgx-a", "all healthy"], ["hgx-b", "healthy"]), "FAIL"),
    # dcgm-G14 (conv 812): the hot GPU called out by its value, without its host
    ("gpu_fleet_health", P, "3 GPUs on 2 hosts. Temperatures are low, with a single RTX A6000 reaching 91 °C – worth "
                            "monitoring for sustained high temps.", "PASS"),
    ("gpu_fleet_health", P, "3 GPUs on 2 hosts. Temperature max 91 °C; all GPUs healthy.", "FAIL"),
    ("gpu_util_avg", P, T(G + ["Avg util %"], ["hgx-a", 0, "2%"], ["hgx-a", 1, "79%"], ["hgx-b", 0, "1%"]), "PASS"),
    ("gpu_util_avg", P, T(["Host", "Avg util %"], ["hgx-a", "41%"], ["hgx-b", "1%"]), "PASS"),          # per host
    ("gpu_util_avg", P, "Fleet average GPU utilization over the last 7 days: 27.5%.", "PASS"),
    ("gpu_util_avg", P, T(G + ["Avg util %"], ["hgx-a", 0, "2%"], ["hgx-a", 1, "40%"]), "FAIL"),
    ("gpu_util_avg", P, T(["Host", "Avg util %"], ["hgx-a", "80%"], ["hgx-b", "1%"]), "FAIL"),
    # prometheus-G15 (conv 910): hosts named by ip; "averaged ≈ 78 %" was the mean of 5-minute peaks -> FAIL
    ("gpu_util_avg", P, "Most GPUs were idle. 10.0.0.1 averaged ≈ 78 % utilization, 10.0.0.9 about 1 %. "
                        "![](ui://line-chart-1)", "FAIL"),
    ("gpu_util_avg", P, "10.0.0.1 averaged ≈ 41 % utilization over the week, 10.0.0.9 about 1 %.", "PASS"),
]
PARAM = {"gpu_util_chart": 1, "power_temp_window": 24, "gpu_idle": 5, "gpu_fleet_health": 85, "gpu_util_avg": 168}
TOL = {"gpu_util_avg": 5}


def run(check, src, answer, trace=()):
    row = PromptRow("GX", "prompt", check, PARAM.get(check), TOL.get(check))
    return evaluate(Ctx(row, src, answer, "[Image" in answer, None, list(trace)))


@pytest.mark.parametrize("check, make, answer, expected", CASES, ids=[f"{c[0]}-{c[3]}-{i}" for i, c in enumerate(CASES)])
def test_case(check, make, answer, expected):
    v = run(check, make(), answer)
    assert v.status == expected, (v.reason, v.expected)


def test_every_gpu_check_has_a_passing_and_a_failing_answer():
    assert set(gpu.CHECKS) <= set(CHECKS)
    assert set(gpu.CHECKS) <= {c for c, _, _, e in CASES if e in ("PASS", "NA")}
    assert set(gpu.CHECKS) <= {c for c, _, _, e in CASES if e == "FAIL"}


def test_chart_tool_call_in_agent_trace_counts_as_a_chart_only_when_rendered():
    shown = [{"tool_name": "show_prometheus_chart", "success": True, "result": '{"chart": "rendered", "series_count": 14}'}]
    empty = [{"tool_name": "show_prometheus_chart", "success": True,
              "result": "No series matched 'count by (severity) (ALERTS)' in the last 5m — nothing to chart."}]
    assert run("gpu_util_chart", FakeProm(), "GPU utilization over the last hour is shown above.", shown).status == "PASS"
    # prometheus-G16 (conv 808): no alerts, the chart tool found nothing, NCP said so -> PASS
    assert run("alerts_chart", FakeProm(), "There are currently no active alerts, so no pie chart.", empty).status == "PASS"


def test_throttling_values_are_read_at_ncps_tool_call_time():
    trace = [{"tool_name": "query_dcgm", "timestamp": "2026-10-08T06:32:06+00:00", "duration_ms": 39000},
             {"tool_name": "get_gpu_telemetry", "timestamp": "2026-10-08T06:31:28+00:00", "duration_ms": 58}]
    row = T(G + ["Temp (°C)", "Util (%)", "Power (W)"], ["hgx-a", 1, 70, 95, 300])
    v = run("throttling", FakeProm(), row, trace)                 # 70 °C at the tool-call time (90 °C now)
    assert v.status == "PASS" and "at NCP's tool-call times" in v.reason, v.reason
    assert run("throttling", FakeProm(), row.replace("| 70 |", "| 90 |"), trace).status == "FAIL"
    from ncp_suite.grading.source_view import source_view
    src = FakeProm()
    run("throttling", src, row, trace)
    view = source_view("throttling", src)                         # the reads that were graded, with SM clock
    assert "SM clock MHz" in view and "| hgx-a GPU 1 | 06:31:28 UTC |" in view and "| 70 |" in view, view


def test_throttling_series_present_names_throttled_gpus():
    src = FakeProm(throttle={"hgx-a:1": 0x40})                      # 0x40 HW slowdown; 0x1 (idle) does not count
    assert run("throttling", src, "hgx-a GPU 1 is throttled (HW slowdown).").status == "PASS"
    assert run("throttling", src, "No GPUs are throttled.").status == "FAIL"


def test_alerts_present_must_be_named_with_severity():
    src = FakeProm(alerts=[{"name": "GPUHot", "severity": "critical"}])
    assert run("alerts_by_severity", src, "critical: GPUHot (hgx-a)").status == "PASS"
    assert run("alerts_by_severity", src, "No alerts are firing.").status == "FAIL"


def test_dcgm_lists_only_its_supported_metrics():
    supported = DcgmSource.catalog(FakeDcgm())                      # data/dcgm_supported_metrics.csv ∩ Prometheus
    assert supported == ["DCGM_FI_DEV_GPU_TEMP", "DCGM_FI_DEV_GPU_UTIL", "DCGM_FI_DEV_POWER_USAGE"]
    with pytest.raises(Unsupported):
        FakeDcgm().alerts()


def test_gpu_reading_helpers():
    assert gpu.indexes("GPUs 0-3 idle") == {"0", "1", "2", "3"} and gpu.indexes("GPU 0, 1 and 2") == {"0", "1", "2"}
    assert [(v, u) for _, v, u in gpu.quantities("92°C, 142.1 W, 47,105 MiB, 50th, GPU1")] == \
        [(92.0, "c"), (142.1, "w"), (47105.0, "mib")]


def test_gpu_source_view_shows_the_compared_values():
    from ncp_suite.grading.source_view import source_view
    assert "| hgx-a GPU 1 | NVIDIA RTX A6000 | 95 | 90 | 300 |" in source_view("gpu_temp_power", FakeProm())
    assert "util avg (168 h)" in source_view("gpu_util_avg", FakeProm(), None, 168)
    assert source_view("alerts_chart", FakeProm()).startswith("No alerts firing")


# ---------------------------------------------------------------- settings, sheet, plugin
def test_gpu_sheet_and_registry():
    rows = load_prompts(settings.GPU_PROMPTS_XLSX)
    assert [r.id for r in rows] == [f"G{i:02d}" for i in range(1, 17)] and all(r.check in CHECKS for r in rows)
    for key in ("prometheus", "dcgm"):
        c = settings.CONNECTORS[key]
        assert (c.prompts, c.needs_login) == (settings.GPU_PROMPTS_XLSX, False)


def test_login_keys_needed_only_where_the_source_logs_in():
    gpu_conn = settings.Connector("demo", "Demo", "#demo", "x:y", "http://x", "", "", needs_login=False)
    assert settings.missing([gpu_conn], ncp=False) == []
    assert settings.missing([settings.Connector("demo", "Demo", "#demo", "x:y", "http://x", "", "")], ncp=False) == \
        ["DEMO_USER", "DEMO_PASSWORD"]


def test_each_connector_reads_its_own_sheet():
    from ncp_suite.pytest_plugin import selected_prompts

    class Config:
        def __init__(self, connectors="", excel=""):
            self.opts = {"--connectors": connectors, "--prompts": "", "--excel": excel}

        def getoption(self, name):
            return self.opts[name]
    rows = {r.id: r for r in selected_prompts(Config())}
    assert len(rows) == 36
    assert rows["P01"].applies_to == {"nexus", "catalyst", "zabbix", "ones"}
    assert rows["G06"].applies_to == {"prometheus", "dcgm"} and not rows["G06"].applies("ones")
    assert [r.id for r in selected_prompts(Config("dcgm"))] == [f"G{i:02d}" for i in range(1, 17)]
    assert len(selected_prompts(Config("dcgm", str(settings.PROMPTS_XLSX)))) == 20      # --excel: one sheet for all


def test_excel_matrix_marks_prompts_of_another_sheet(tmp_path):
    from openpyxl import load_workbook

    from ncp_suite.pytest_plugin import selected_prompts
    from ncp_suite.reporting.excel import write_report
    from ncp_suite.results import PromptResult

    class Config:
        def getoption(self, name):
            return {"--connectors": "ones,dcgm", "--prompts": "P01,G06", "--excel": ""}[name]
    prompts = selected_prompts(Config())
    conns = [settings.CONNECTORS["ones"], settings.CONNECTORS["dcgm"]]
    results = [PromptResult("P01", "ones", "ONES", "PASS"), PromptResult("G06", "dcgm", "DCGM", "FAIL", "x")]
    ws = load_workbook(write_report(results, prompts, conns, tmp_path))["Matrix"]
    assert [[c.value for c in r][:3] for r in ws.iter_rows(min_row=2)] == \
        [["List all devices in my network.", "PASS", "—"], ["Show current GPU temperature and power draw.", "—", "FAIL"]]


def test_fleet_summary_units_are_not_a_gpu_count():
    # prometheus-G14 (conv 800): "> 200 W per GPU" is not "200 GPUs"; the hot GPU must be flagged
    v = run("gpu_fleet_health", FakeProm(), "Average power 89.8 W, well below typical maximums (often > 200 W per "
                                            "GPU). GPU count: 14. Overall the fleet looks healthy.")
    assert v.status == "FAIL" and v.reason.startswith("not flagged: hgx-a GPU 1"), v.reason
