"""Excel report: the same matrix as Dev's sheet, plus details and a summary.

Based on: report.py (2026-10-06). CHANGED: reads PromptResult; Details rows sorted by prompt
then connector (parallel workers finish in any order); new columns "Tools called", "Retries" and
"Source data" (the source table the check compared, next to NCP's answer).
From Automation 2 (USECASE / Ticketing conftest.py): control characters are removed before a
value is written (openpyxl refuses them, which lost the whole workbook), and the FAILs get their
own sheet (the old "_failed_prompts.xlsx" file, here a sheet in the same workbook).
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ncp_suite.results import FILL, LEGEND, STATUSES, PromptResult, by_cell, worst

FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
HEAD = Font(name="Arial", size=11, bold=True)
HEAD_COLOUR = "B4A7D6"                                         # Dev's sheet header colour
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
EDGE = Side(style="thin", color="000000")
BOX = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)


def _cell(ws, row, col, value, font=FONT, align=WRAP, fill=None):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)       # NCP answers can carry control characters
    c = ws.cell(row=row, column=col, value=value)
    c.font, c.alignment, c.border = font, align, BOX
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c


def _head(ws, col, value):
    return _cell(ws, 1, col, value, HEAD, CENTER, HEAD_COLOUR)


def write_report(results: list[PromptResult], prompts: list, connectors: list, out_dir: Path,
                 stamp: str | None = None) -> Path:
    """stamp = the run's time stamp, so the .xlsx and the .html of one run share a name."""
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"NCP_MCP_Prompt_Results_{stamp or f'{datetime.now():%Y%m%d_%H%M%S}'}.xlsx"
    order = {(p.id, c.key): (i, j) for i, p in enumerate(prompts) for j, c in enumerate(connectors)}
    results = sorted(results, key=lambda r: order.get((r.id, r.connector), (len(order), 0)))
    wb = Workbook()

    # ---- Matrix: Prompt | <connector titles> | Comments -------------------------------------
    ws = wb.active
    ws.title = "Matrix"
    headers = ["Prompt"] + [c.title for c in connectors] + ["Comments"]
    for i, h in enumerate(headers, 1):
        _head(ws, i, h)
    ws.column_dimensions["A"].width = 62
    for i in range(2, len(headers)):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 22
    ws.column_dimensions[ws.cell(row=1, column=len(headers)).column_letter].width = 80
    cells = by_cell(results)
    for row_i, p in enumerate(prompts, 2):
        _cell(ws, row_i, 1, p.prompt)
        comments = []
        for col_i, conn in enumerate(connectors, 2):
            runs = cells.get((p.id, conn.key), [])
            if not p.applies(conn.key):
                _cell(ws, row_i, col_i, "—", align=CENTER)            # not in this connector's prompt sheet
                continue
            if not runs:
                _cell(ws, row_i, col_i, "", align=CENTER)
                continue
            status = worst(r.status for r in runs)
            _cell(ws, row_i, col_i, status, BOLD, CENTER, FILL.get(status))
            for r in runs:
                if r.status != "PASS":
                    scope = f" [{r.scope}]" if len(runs) > 1 else ""
                    comments.append(f"{conn.title.split(' (')[0]}{scope}: {r.status} — {r.reason}")
        _cell(ws, row_i, len(headers), "\n".join(comments))
    ws.freeze_panes = "B2"

    # ---- Details -------------------------------------------------------------------------------
    d = wb.create_sheet("Details")
    cols = ["ID", "Connector", "Scope", "Result", "Reason", "Expected (source)", "Prompt sent", "Device",
            "Follow-ups", "Seconds", "Conversation", "Tools called", "Retries", "Judge note", "Source data",
            "NCP answer", "NCP tool calls (agent_trace)", "LLM reader table"]
    widths = [6, 24, 8, 9, 60, 60, 50, 18, 40, 8, 12, 30, 40, 40, 90, 90, 90, 60]
    _rows(d, cols, widths, [_details_row(r) for r in results], status_col=4)

    # ---- Failures: what to look at first (FAIL and XFAIL only) -------------------------------------
    f = wb.create_sheet("Failures")
    failed = [r for r in results if r.status in ("FAIL", "XFAIL")]
    _rows(f, ["ID", "Connector", "Result", "Reason", "Expected (source)", "Conversation", "Tools called",
              "Retries", "Follow-ups", "Prompt sent", "Source data", "NCP answer", "NCP tool calls (agent_trace)",
              "LLM reader table"],
          [6, 24, 9, 60, 60, 12, 30, 40, 40, 50, 90, 90, 90, 60],
          [[r.id, r.title, r.status, r.reason, r.expected, r.conversation_id, ", ".join(r.tools),
            "\n".join(r.retries), _followups(r), r.sent, r.source[:32000], (r.answer or "")[:32000],
            _trace(r), r.extracted[:32000]]
           for r in failed], status_col=3)

    # ---- Summary (live formulas over Details) --------------------------------------------------
    s = wb.create_sheet("Summary", 0)
    _head(s, 1, "Connector")
    for j, st in enumerate([*STATUSES, "Total"], 2):
        _head(s, j, st)
    last = len(results) + 1
    for i, conn in enumerate(connectors, 2):
        _cell(s, i, 1, conn.title, BOLD)
        for j, st in enumerate(STATUSES, 2):
            _cell(s, i, j, f'=COUNTIFS(Details!$B$2:$B${last},$A{i},Details!$D$2:$D${last},"{st}")',
                  align=CENTER)
        _cell(s, i, len(STATUSES) + 2, f"=SUM(B{i}:F{i})", BOLD, CENTER)
    s.column_dimensions["A"].width = 30
    s.cell(row=len(connectors) + 3, column=1, value=LEGEND).font = FONT
    wb.save(path)
    return path


def _followups(r: PromptResult) -> str:
    return " || ".join(f"{a} -> {b[:150]}" for a, b in r.followups)


def _trace(r: PromptResult) -> str:
    """agent_trace for a cell: the summary, then one line per call (results already shortened)."""
    lines = [r.trace_summary] + [f"{c.get('path')} | {c.get('tool')} | ok={c.get('success')} | "
                                 f"{(c.get('error') or c.get('result') or '')[:400]}" for c in r.trace]
    return "\n".join(x for x in lines if x)[:32000]


def _details_row(r: PromptResult) -> list:
    return [r.id, r.title, r.scope, r.status, r.reason, r.expected, r.sent, r.device, _followups(r), r.seconds,
            r.conversation_id, ", ".join(r.tools), "\n".join(r.retries), r.judge, r.source[:32000],
            (r.answer or "")[:32000], _trace(r), r.extracted[:32000]]


def _rows(ws, cols: list[str], widths: list[int], rows: list[list], status_col: int) -> None:
    """Header, column widths, rows (status cell coloured), frozen first columns, filter."""
    for i, (h, w) in enumerate(zip(cols, widths), 1):
        _head(ws, i, h)
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for i, values in enumerate(rows, 2):
        for j, v in enumerate(values, 1):
            _cell(ws, i, j, v, fill=FILL.get(values[status_col - 1]) if j == status_col else None)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(cols)).column_letter}{max(2, len(rows) + 1)}"
